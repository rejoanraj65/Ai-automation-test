"""
Single-file AI Price List Automation MVP
========================================
Run locally:
    pip install fastapi uvicorn requests openpyxl apscheduler python-multipart python-docx
    uvicorn app:app --host 0.0.0.0 --port 8000

Render:
    Build command: pip install -r requirements.txt
    Start command: uvicorn app:app --host 0.0.0.0 --port $PORT

IMPORTANT:
- This prototype intentionally has NO login, as requested. Anyone who can open the
  website can create/edit/delete client profiles and change the Gemini API key.
  Do not put real secrets or sensitive client data here until access control is added.
- Render Free has an ephemeral filesystem and sleeps. SQLite data and generated files
  may disappear on restart/redeploy; hourly jobs do not run while the service sleeps.
- Public data endpoints are intentionally readable by anyone.
"""

import os
import re
import io
import json
import time
import uuid
import sqlite3
import logging
from datetime import datetime, timezone
from urllib.parse import urlparse

import requests
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse, PlainTextResponse
from fastapi.middleware.cors import CORSMiddleware
from apscheduler.schedulers.background import BackgroundScheduler
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

try:
    from docx import Document
except Exception:
    Document = None

APP_TITLE = "PriceList AI"
DB_PATH = os.getenv("DATABASE_PATH", "pricelist.db")
POLL_MINUTES = max(5, int(os.getenv("POLL_MINUTES", "60")))
GEMINI_API_BASE = "https://generativelanguage.googleapis.com/v1beta"
MAX_UPLOAD_BYTES = 15 * 1024 * 1024
HTTP_TIMEOUT = 35

logging.basicConfig(level=logging.INFO)
log = logging.getLogger("pricelist-ai")

app = FastAPI(title=APP_TITLE, version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


def db():
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with db() as con:
        con.execute("""
            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            )
        """)
        con.execute("""
            CREATE TABLE IF NOT EXISTS clients (
                id TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                source_url TEXT DEFAULT '',
                source_type TEXT DEFAULT 'text',
                source_text TEXT DEFAULT '',
                products_json TEXT DEFAULT '[]',
                source_hash TEXT DEFAULT '',
                last_checked TEXT DEFAULT '',
                last_updated TEXT DEFAULT '',
                status TEXT DEFAULT 'Waiting',
                last_error TEXT DEFAULT '',
                created_at TEXT NOT NULL
            )
        """)
        con.execute("""
            CREATE TABLE IF NOT EXISTS history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                client_id TEXT NOT NULL,
                event TEXT NOT NULL,
                details TEXT DEFAULT '',
                created_at TEXT NOT NULL
            )
        """)
        con.commit()


def now_iso():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def setting_get(key, default=""):
    with db() as con:
        row = con.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
        return row["value"] if row else default


def setting_set(key, value):
    with db() as con:
        con.execute(
            "INSERT INTO settings(key,value) VALUES(?,?) "
            "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (key, value),
        )
        con.commit()


def log_event(client_id, event, details=""):
    with db() as con:
        con.execute(
            "INSERT INTO history(client_id,event,details,created_at) VALUES(?,?,?,?)",
            (client_id, event, str(details)[:2000], now_iso()),
        )
        con.commit()


def slugify(value):
    value = value.strip().lower()
    value = re.sub(r"[^a-z0-9]+", "-", value).strip("-")
    return value[:45] or "client"


def make_client_id(name):
    base = slugify(name)
    return f"{base}-{uuid.uuid4().hex[:6]}"


def get_client(client_id):
    with db() as con:
        row = con.execute("SELECT * FROM clients WHERE id=?", (client_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Client not found")
    return dict(row)


def update_client(client_id, **fields):
    allowed = {
        "name", "source_url", "source_type", "source_text", "products_json",
        "source_hash", "last_checked", "last_updated", "status", "last_error"
    }
    fields = {k: v for k, v in fields.items() if k in allowed}
    if not fields:
        return
    sql = ", ".join(f"{k}=?" for k in fields)
    values = list(fields.values()) + [client_id]
    with db() as con:
        con.execute(f"UPDATE clients SET {sql} WHERE id=?", values)
        con.commit()


def sha256_text(text):
    import hashlib
    return hashlib.sha256(text.encode("utf-8", errors="ignore")).hexdigest()


def extract_google_doc_id(url):
    patterns = [
        r"/document/d/([a-zA-Z0-9_-]+)",
        r"[?&]id=([a-zA-Z0-9_-]+)",
    ]
    for pattern in patterns:
        m = re.search(pattern, url)
        if m:
            return m.group(1)
    raise ValueError("Google Docs URL থেকে Document ID পাওয়া যায়নি।")


def fetch_public_google_doc(url):
    doc_id = extract_google_doc_id(url)
    export_url = f"https://docs.google.com/document/d/{doc_id}/export?format=txt"
    response = requests.get(
        export_url,
        timeout=HTTP_TIMEOUT,
        headers={"User-Agent": "Mozilla/5.0 PriceListAI/1.0"},
    )
    if response.status_code != 200:
        raise ValueError(
            "Google Doc পড়া যায়নি। Sharing সেটিংস 'Anyone with the link – Viewer' "
            "করো এবং লিংকটি ঠিক আছে কি না দেখো।"
        )
    text = response.content.decode("utf-8", errors="replace").strip()
    if not text or "accounts.google.com" in text[:1000]:
        raise ValueError("ডকুমেন্টের টেক্সট পাওয়া যায়নি। ডকুমেন্টটি Public Viewer কি না যাচাই করো।")
    return text


def read_uploaded_file(filename, raw):
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext in ("txt", "csv", "md"):
        return raw.decode("utf-8-sig", errors="replace")
    if ext == "docx":
        if Document is None:
            raise ValueError("DOCX পড়তে python-docx ইনস্টল করতে হবে।")
        document = Document(io.BytesIO(raw))
        lines = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(lines)
    if ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if vals:
                    lines.append(" | ".join(vals))
        return "\n".join(lines)
    if ext == "xls":
        raise ValueError("পুরোনো .xls নয়; Excel ফাইলটি .xlsx হিসেবে Save করে আপলোড করো।")
    raise ValueError("সমর্থিত ফাইল: TXT, CSV, MD, DOCX, XLSX")


def discover_gemini_models(api_key):
    response = requests.get(
        f"{GEMINI_API_BASE}/models",
        params={"key": api_key},
        timeout=HTTP_TIMEOUT,
    )
    if response.status_code != 200:
        raise ValueError(f"Gemini model list error: {response.status_code} {response.text[:400]}")
    models = response.json().get("models", [])
    usable = []
    for model in models:
        methods = model.get("supportedGenerationMethods", [])
        name = model.get("name", "")
        if "generateContent" in methods and name:
            short_name = name.replace("models/", "")
            usable.append({
                "name": short_name,
                "display_name": model.get("displayName", short_name),
                "input_token_limit": model.get("inputTokenLimit"),
                "output_token_limit": model.get("outputTokenLimit"),
            })
    return usable


def order_models(models):
    # Prefer Flash for routine extraction; keep the remaining compatible models
    # as fallbacks if a model is overloaded, unavailable, or rejects the request.
    def rank(model):
        name = model["name"].lower()
        if "flash" in name:
            return (0, name)
        if "pro" in name:
            return (1, name)
        return (2, name)
    ordered = sorted(models, key=rank)
    if not ordered:
        raise ValueError("এই API Key-তে generateContent সমর্থিত কোনো Gemini model পাওয়া যায়নি।")
    return ordered


def call_gemini(api_key, model, source_text):
    prompt = """
You are a precise product price-list extraction engine.
Extract every product from the source text. The source may contain Bangla or English,
brand headings, messy line breaks, spelling errors, and inconsistent formatting.

Return ONLY valid JSON, no markdown, using exactly this shape:
{
  "products": [
    {
      "brand": "Brand name or empty string",
      "model": "Full product/model name",
      "variant": "RAM/ROM or other variant, empty if absent",
      "price": 0,
      "currency": "BDT",
      "notes": ""
    }
  ],
  "warnings": ["uncertain or ambiguous items"]
}

Rules:
- Never invent a price or product.
- price must be a number when clearly present; otherwise null.
- Keep distinct RAM/ROM variants as separate products.
- Do not merge two similar but different models.
- Ignore headings, phone numbers, dates, and unrelated text.
- Preserve Bangla product/model text where appropriate.
- If a line is ambiguous, include it only when reasonably identifiable and add a warning.
- Do not treat an old price as current unless the source says so.
SOURCE DOCUMENT:
""" + source_text[:180000]

    url = f"{GEMINI_API_BASE}/models/{model}:generateContent"
    response = requests.post(
        url,
        params={"key": api_key},
        json={
            "contents": [{"parts": [{"text": prompt}]}],
            "generationConfig": {
                "temperature": 0.1,
                "responseMimeType": "application/json",
            },
        },
        timeout=90,
    )
    if response.status_code != 200:
        raise ValueError(f"Gemini API error ({response.status_code}): {response.text[:700]}")
    payload = response.json()
    try:
        answer = payload["candidates"][0]["content"]["parts"][0]["text"]
    except (KeyError, IndexError, TypeError):
        raise ValueError("Gemini থেকে ব্যবহারযোগ্য response পাওয়া যায়নি।")
    answer = answer.strip()
    answer = re.sub(r"^```(?:json)?\s*", "", answer)
    answer = re.sub(r"\s*```$", "", answer)
    data = json.loads(answer)
    products = normalize_products(data.get("products", []))
    return products, data.get("warnings", [])


def normalize_products(items):
    cleaned = []
    seen = set()
    for item in items:
        if not isinstance(item, dict):
            continue
        brand = str(item.get("brand") or "").strip()
        model = str(item.get("model") or "").strip()
        variant = str(item.get("variant") or "").strip()
        notes = str(item.get("notes") or "").strip()
        price = item.get("price")
        if not model:
            continue
        try:
            if price is not None and price != "":
                price = float(str(price).replace(",", "").replace("৳", "").strip())
                if price.is_integer():
                    price = int(price)
            else:
                price = None
        except Exception:
            price = None
        key = (brand.casefold(), model.casefold(), variant.casefold())
        if key in seen:
            continue
        seen.add(key)
        cleaned.append({
            "brand": brand,
            "model": model,
            "variant": variant,
            "price": price,
            "currency": str(item.get("currency") or "BDT"),
            "notes": notes,
        })
    return cleaned


def process_client(client_id, force=False):
    client = get_client(client_id)
    api_key = setting_get("gemini_api_key")
    if not api_key:
        update_client(client_id, status="Needs API Key", last_error="Gemini API Key সেট করা হয়নি।")
        return {"ok": False, "error": "Gemini API Key সেট করা হয়নি।"}

    try:
        if client["source_type"] == "google_docs":
            source_text = fetch_public_google_doc(client["source_url"])
        else:
            source_text = client["source_text"]

        if not source_text.strip():
            raise ValueError("Source document খালি।")

        current_hash = sha256_text(source_text)
        if not force and current_hash == client.get("source_hash") and client.get("products_json"):
            update_client(client_id, last_checked=now_iso(), status="No Changes", last_error="")
            log_event(client_id, "Hourly check", "কোনো পরিবর্তন পাওয়া যায়নি।")
            return {"ok": True, "changed": False, "message": "কোনো পরিবর্তন নেই।"}

        models = order_models(discover_gemini_models(api_key))
        products, warnings = None, []
        model, model_errors = "", []
        for candidate in models:
            try:
                products, warnings = call_gemini(api_key, candidate["name"], source_text)
                model = candidate["name"]
                break
            except Exception as model_exc:
                model_errors.append(f"{candidate['name']}: {model_exc}")
                log.warning("Gemini model fallback: %s", model_errors[-1])
        if products is None:
            raise ValueError("সব উপলব্ধ Gemini model ব্যর্থ হয়েছে: " + " | ".join(model_errors)[:1500])

        if not products:
            raise ValueError("AI কোনো পণ্য শনাক্ত করেনি; আগের ডেটা নিরাপদ রাখতে আপডেট বাতিল করা হয়েছে।")

        old_products = json.loads(client.get("products_json") or "[]")
        old_map = {
            (p.get("brand", "").casefold(), p.get("model", "").casefold(), p.get("variant", "").casefold()): p
            for p in old_products
        }
        new_map = {
            (p.get("brand", "").casefold(), p.get("model", "").casefold(), p.get("variant", "").casefold()): p
            for p in products
        }
        added = len(set(new_map) - set(old_map))
        removed = len(set(old_map) - set(new_map))
        changed = sum(1 for k in set(new_map) & set(old_map) if new_map[k] != old_map[k])

        update_client(
            client_id,
            source_text=source_text,
            products_json=json.dumps(products, ensure_ascii=False),
            source_hash=current_hash,
            last_checked=now_iso(),
            last_updated=now_iso(),
            status="Processed",
            last_error="",
        )
        details = f"Model: {model}; Products: {len(products)}; Added: {added}; Changed: {changed}; Removed: {removed}; Warnings: {len(warnings)}"
        log_event(client_id, "Processed", details)
        if warnings:
            log_event(client_id, "AI warnings", json.dumps(warnings, ensure_ascii=False)[:1800])
        return {
            "ok": True, "changed": True, "model": model, "total_products": len(products),
            "added": added, "changed_products": changed, "removed": removed, "warnings": warnings,
        }
    except Exception as exc:
        log.exception("Processing failed for %s", client_id)
        update_client(client_id, status="Error", last_checked=now_iso(), last_error=str(exc))
        log_event(client_id, "Error", str(exc))
        return {"ok": False, "error": str(exc)}


def hourly_sync():
    try:
        with db() as con:
            rows = con.execute(
                "SELECT id FROM clients WHERE source_type='google_docs' AND source_url<>''"
            ).fetchall()
        for row in rows:
            process_client(row["id"], force=False)
            time.sleep(1)
    except Exception:
        log.exception("Hourly sync failed")


def build_excel(products, client_name):
    wb = Workbook()
    ws = wb.active
    ws.title = "Price List"
    headers = ["Brand", "Model", "Variant", "Price", "Currency", "Notes"]
    ws.append([client_name])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.append(headers)
    for cell in ws[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2874A6")
        cell.alignment = Alignment(horizontal="center")
    for p in products:
        ws.append([
            p.get("brand"), p.get("model"), p.get("variant"),
            p.get("price"), p.get("currency"), p.get("notes")
        ])
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:F{max(2, ws.max_row)}"
    widths = [20, 34, 20, 15, 12, 35]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def public_base_url(request=None):
    configured = setting_get("public_base_url").strip().rstrip("/")
    if configured:
        return configured
    if request:
        return str(request.base_url).rstrip("/")
    return ""


@app.on_event("startup")
def startup():
    init_db()
    scheduler = BackgroundScheduler(daemon=True, timezone="UTC")
    scheduler.add_job(hourly_sync, "interval", minutes=POLL_MINUTES, id="hourly_sync", replace_existing=True)
    scheduler.start()
    app.state.scheduler = scheduler
    log.info("Started. Hourly sync interval: %s minutes", POLL_MINUTES)


@app.on_event("shutdown")
def shutdown():
    scheduler = getattr(app.state, "scheduler", None)
    if scheduler:
        scheduler.shutdown(wait=False)


def esc_html(value):
    return (str(value if value is not None else "")
            .replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;").replace('"', "&quot;").replace("'", "&#39;"))


def render_page(title, active, content, extra_script=""):
    nav_items = [
        ("/dashboard", "Dashboard", "dashboard"),
        ("/clients", "Clients", "clients"),
        ("/clients/new", "Add Client", "new"),
        ("/settings", "Settings", "settings"),
        ("/api-guide", "API Guide", "api"),
    ]
    nav = "".join(
        f'<a class="nav-link {"active" if key == active else ""}" href="{url}">{label}</a>'
        for url, label, key in nav_items
    )
    return HTMLResponse(f'''<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>{esc_html(title)} · {APP_TITLE}</title>
<style>
:root{{--bg:#f4f6fb;--card:#fff;--ink:#172033;--muted:#667085;--line:#e5eaf2;--blue:#315efb;--navy:#13213b;--green:#087443}}
*{{box-sizing:border-box}}body{{margin:0;background:var(--bg);color:var(--ink);font:15px Inter,Arial,sans-serif}}
.topbar{{background:linear-gradient(120deg,#13213b,#244b87);color:white;padding:20px max(20px,calc((100% - 1180px)/2))}}
.brand{{font-weight:800;font-size:21px;letter-spacing:.2px}}.sub{{color:#d5e1f5;font-size:12px;margin-top:4px}}
.nav{{background:#fff;border-bottom:1px solid var(--line);display:flex;gap:5px;align-items:center;padding:0 max(12px,calc((100% - 1180px)/2));overflow-x:auto;white-space:nowrap}}
.nav-link{{display:inline-block;padding:15px 13px;color:#475467;text-decoration:none;font-weight:650;font-size:13px;border-bottom:3px solid transparent}}
.nav-link:hover{{color:var(--blue);background:#f7f9ff}}.nav-link.active{{color:var(--blue);border-bottom-color:var(--blue)}}
main{{max-width:1180px;margin:26px auto;padding:0 16px;min-height:65vh}}h1{{font-size:25px;margin:0 0 8px}}h2{{font-size:17px;margin:0 0 14px}}p{{line-height:1.55}}.muted,.small{{color:var(--muted)}}.small{{font-size:12px}}
.grid{{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:16px}}.grid-3{{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:14px}}
.card{{background:var(--card);border:1px solid var(--line);border-radius:15px;padding:18px;box-shadow:0 6px 24px #17203308;margin-bottom:16px}}
.stat{{font-size:28px;font-weight:800;margin-top:8px}}.field{{margin:13px 0}}label{{display:block;font-size:13px;color:#475467;font-weight:650;margin-bottom:6px}}
input,select,textarea{{width:100%;border:1px solid #d6deea;border-radius:9px;padding:11px 12px;font:inherit;background:#fff;color:var(--ink)}}textarea{{min-height:90px;resize:vertical}}
button,.btn{{display:inline-block;border:0;border-radius:9px;background:var(--blue);color:white;padding:11px 15px;font-weight:650;cursor:pointer;text-decoration:none;font-size:13px}}
button.secondary,.btn.secondary{{background:#eaf0ff;color:#2347c8}}button.danger,.btn.danger{{background:#fff0f0;color:#b42318}}button:disabled{{opacity:.55;cursor:wait}}
.row{{display:flex;gap:9px;align-items:center;flex-wrap:wrap}}.between{{justify-content:space-between}}.notice{{padding:12px 14px;border-radius:10px;background:#fff7e6;color:#875a00;font-size:12px;line-height:1.5;margin-bottom:16px}}
.client{{border:1px solid var(--line);border-radius:12px;padding:14px;margin:10px 0;background:#fff}}.client strong{{font-size:15px}}.status{{font-size:11px;border-radius:20px;padding:5px 9px;background:#edf2f7;color:#475467}}
a{{color:var(--blue);text-decoration:none;word-break:break-word}}.actions{{display:flex;gap:7px;flex-wrap:wrap;margin-top:12px}}.actions button{{padding:8px 10px;font-size:12px}}
.table-wrap{{overflow:auto;background:#fff;border:1px solid var(--line);border-radius:12px}}table{{border-collapse:collapse;width:100%;min-width:620px}}th,td{{padding:11px 12px;border-bottom:1px solid var(--line);text-align:left;font-size:13px}}th{{background:#f0f4fa}}
#message,.message{{margin:12px 0;color:#2347c8;font-size:13px;white-space:pre-wrap}}code,pre{{background:#f3f5f9;border-radius:7px}}code{{padding:2px 5px}}pre{{padding:13px;overflow:auto;font-size:12px;line-height:1.55}}
footer{{max-width:1180px;margin:12px auto 28px;padding:0 16px;color:var(--muted);font-size:12px}}
@media(max-width:760px){{.grid,.grid-3{{grid-template-columns:1fr}}.topbar{{padding:18px}}main{{margin:18px auto}}}}
</style></head><body>
<header class="topbar"><div class="brand">PriceList AI</div><div class="sub">Product data management · AI-ready search · Public price lists</div></header>
<nav class="nav">{nav}</nav><main>{content}</main><footer>PriceList AI · Multi-page management dashboard · Public product data remains separately accessible.</footer>
<script>async function api(url,opts={{}}){{let r=await fetch(url,opts);let d=await r.json().catch(()=>({{detail:"Unexpected server response"}}));if(!r.ok)throw new Error(d.detail||d.error||"Request failed");return d}}
function esc(s){{return String(s??'').replace(/[&<>"']/g,c=>({{'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}}[c]))}}
function showMessage(t){{let e=document.getElementById('message');if(e)e.textContent=t;}}
{extra_script}</script></body></html>''')


@app.get("/", response_class=HTMLResponse)
@app.get("/dashboard", response_class=HTMLResponse)
def home():
    with db() as con:
        rows = con.execute("SELECT status, products_json FROM clients").fetchall()
    total_clients = len(rows)
    total_products = sum(len(json.loads(r["products_json"] or "[]")) for r in rows)
    errors = sum(1 for r in rows if r["status"] == "Error")
    content = f'''<h1>Dashboard</h1><p class="muted">তোমার Price List সিস্টেমের সার্বিক অবস্থা এক জায়গায় দেখো।</p>
<div class="notice"><b>Prototype security:</b> এখনো Login/Authentication নেই। Dashboard-এর URL জানা যে কেউ settings ও client data পরিবর্তন করতে পারে। বাস্তব sensitive data ব্যবহারের আগে authentication যোগ করো। Render Free-তে local SQLite data restart/redeploy হলে হারাতে পারে।</div>
<div class="grid-3"><section class="card"><div class="muted">Total Clients</div><div class="stat">{total_clients}</div></section>
<section class="card"><div class="muted">Extracted Products</div><div class="stat">{total_products}</div></section>
<section class="card"><div class="muted">Clients with Errors</div><div class="stat">{errors}</div></section></div>
<div class="grid"><section class="card"><h2>Quick Actions</h2><p><a class="btn" href="/clients/new">＋ Add New Client</a></p><p><a class="btn secondary" href="/clients">Manage Clients</a></p><p><a class="btn secondary" href="/settings">Configure Gemini API</a></p></section>
<section class="card"><h2>System Modules</h2><p>• Client profiles and source documents</p><p>• Gemini extraction and model discovery</p><p>• Hourly Google Docs change detection</p><p>• Public HTML, JSON API and Excel export</p><p>• AI-ready search and filtering endpoints</p></section></div>'''
    return render_page("Dashboard", "dashboard", content)


@app.get("/clients", response_class=HTMLResponse)
def clients_page():
    with db() as con:
        rows = con.execute("SELECT * FROM clients ORDER BY created_at DESC").fetchall()
    cards = []
    for row in rows:
        c = dict(row); count = len(json.loads(c.get("products_json") or "[]"))
        cards.append(f'''<article class="client"><div class="row between"><strong>{esc_html(c['name'])}</strong><span class="status">{esc_html(c['status'])}</span></div>
<div class="small">ID: {esc_html(c['id'])} · {count} products · Last updated: {esc_html(c.get('last_updated') or 'Never')}</div>
<div class="actions"><a class="btn secondary" href="/clients/{esc_html(c['id'])}/manage">Manage Client</a><a class="btn secondary" href="/data/{esc_html(c['id'])}" target="_blank">Public Price List</a><a class="btn secondary" href="/api/public/{esc_html(c['id'])}" target="_blank">JSON Data</a></div>
{('<p class="small" style="color:#b42318">'+esc_html(c.get('last_error'))+'</p>') if c.get('last_error') else ''}</article>''')
    body = "".join(cards) or '<div class="card"><p>No clients yet. Create your first client profile.</p><a class="btn" href="/clients/new">Create Client</a></div>'
    content = f'''<div class="row between"><div><h1>Clients</h1><p class="muted">প্রতিটি Client-এর source, products, public URL ও processing নিয়ন্ত্রণ করো।</p></div><a class="btn" href="/clients/new">＋ Add Client</a></div>{body}'''
    return render_page("Clients", "clients", content)


@app.get("/clients/new", response_class=HTMLResponse)
def new_client_page():
    content = '''<h1>Create Client</h1><p class="muted">একটি নতুন client profile তৈরি করো। Google Docs URL অথবা পরে upload করা file ব্যবহার করতে পারবে।</p>
<section class="card" style="max-width:720px"><div class="field"><label>Client / Shop Name</label><input id="clientName" placeholder="Example: ABC Mobile Shop"></div>
<div class="field"><label>Public Google Docs URL (optional)</label><input id="docUrl" placeholder="https://docs.google.com/document/d/..."></div>
<p class="small">Google Docs sharing অবশ্যই Anyone with the link → Viewer হতে হবে।</p><div class="row"><button onclick="createClient()">Create Profile</button><a class="btn secondary" href="/clients">Cancel</a></div><div id="message"></div></section>'''
    script = '''async function createClient(){try{let name=document.getElementById('clientName').value.trim(),url=document.getElementById('docUrl').value.trim();let d=await api('/api/clients',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({name,source_url:url})});location.href='/clients/'+encodeURIComponent(d.id)+'/manage'}catch(e){showMessage(e.message)}}'''
    return render_page("Create Client", "new", content, script)


@app.get("/clients/{client_id}/manage", response_class=HTMLResponse)
def manage_client_page(client_id: str):
    c = get_client(client_id)
    products = json.loads(c.get("products_json") or "[]")
    with db() as con:
        history = con.execute("SELECT event,details,created_at FROM history WHERE client_id=? ORDER BY id DESC LIMIT 15", (client_id,)).fetchall()
    hist = "".join(f'<tr><td>{esc_html(r["created_at"])}</td><td>{esc_html(r["event"])}</td><td>{esc_html(r["details"])}</td></tr>' for r in history)
    content = f'''<div class="row between"><div><h1>{esc_html(c['name'])}</h1><p class="muted">Client ID: {esc_html(client_id)}</p></div><a class="btn secondary" href="/clients">← All Clients</a></div>
<div class="grid"><section class="card"><h2>Source & Processing</h2><p class="small">Source type: {esc_html(c['source_type'])} · Status: {esc_html(c['status'])}</p>
<div class="field"><label>Google Docs URL</label><input id="docUrl" value="{esc_html(c.get('source_url') or '')}" placeholder="Public Google Docs URL"></div>
<div class="row"><button onclick="saveDocUrl()">Save Source URL</button><button class="secondary" onclick="runProcess('check')">Check Changes</button><button onclick="runProcess('process')">Process Now</button></div>
<hr style="border:0;border-top:1px solid #e5eaf2;margin:18px 0"><div class="field"><label>Upload TXT / CSV / MD / DOCX / XLSX (max 15 MB)</label><input id="sourceFile" type="file" accept=".txt,.csv,.md,.docx,.xlsx"></div><button onclick="uploadFile()">Upload Source File</button><div id="message"></div></section>
<section class="card"><h2>Client Data Links</h2><p><a target="_blank" href="/data/{esc_html(client_id)}">Public HTML Price List</a></p><p><a target="_blank" href="/api/public/{esc_html(client_id)}">JSON API (all products)</a></p><p><a target="_blank" href="/api/public/{esc_html(client_id)}/search?q=">AI Search API</a></p><p><a href="/download/{esc_html(client_id)}.xlsx">Download Excel</a></p><p class="small">Products currently extracted: <b>{len(products)}</b><br>Last checked: {esc_html(c.get('last_checked') or 'Never')}<br>Last updated: {esc_html(c.get('last_updated') or 'Never')}</p></section></div>
<section class="card"><div class="row between"><h2>Recent Activity</h2><button class="danger" onclick="deleteClient()">Delete Client</button></div><div class="table-wrap"><table><thead><tr><th>Date</th><th>Event</th><th>Details</th></tr></thead><tbody>{hist or '<tr><td colspan="3">No activity yet.</td></tr>'}</tbody></table></div></section>'''
    script = f'''const CLIENT_ID={json.dumps(client_id)};
async function saveDocUrl(){{try{{let u=document.getElementById('docUrl').value.trim();let d=await api('/api/clients/'+CLIENT_ID+'/source',{{method:'POST',headers:{{'Content-Type':'application/json'}},body:JSON.stringify({{source_url:u}})}});showMessage(d.message||'Source URL saved.')}}catch(e){{showMessage(e.message)}}}}
async function uploadFile(){{try{{let f=document.getElementById('sourceFile').files[0];if(!f)throw new Error('একটি ফাইল নির্বাচন করো।');let fd=new FormData();fd.append('file',f);let d=await api('/api/clients/'+CLIENT_ID+'/upload',{{method:'POST',body:fd}});showMessage(d.message);}}catch(e){{showMessage(e.message)}}}}
async function runProcess(a){{try{{showMessage('Processing...');let d=await api('/api/clients/'+CLIENT_ID+'/'+a,{{method:'POST'}});showMessage(JSON.stringify(d,null,2));setTimeout(()=>location.reload(),1200)}}catch(e){{showMessage(e.message)}}}}
async function deleteClient(){{if(!confirm('এই client এবং তার history মুছে ফেলবে?'))return;try{{await api('/api/clients/'+CLIENT_ID,{{method:'DELETE'}});location.href='/clients'}}catch(e){{showMessage(e.message)}}}}'''
    return render_page("Manage Client", "clients", content, script)


@app.get("/settings", response_class=HTMLResponse)
def settings_page():
    content = '''<h1>Settings</h1><p class="muted">Gemini API এবং public website configuration।</p><section class="card" style="max-width:760px">
<div class="field"><label>Gemini API Key</label><input id="apiKey" type="password" placeholder="Loading settings..."></div>
<div class="field"><label>Public Website Base URL (optional)</label><input id="baseUrl" placeholder="https://your-app.onrender.com"></div>
<div class="row"><button onclick="saveSettings()">Save Settings</button><button class="secondary" onclick="loadModels()">Discover Gemini Models</button></div><div id="models" class="small" style="margin-top:12px"></div><div id="message"></div>
</section>'''
    script = '''async function initSettings(){try{let s=await api('/api/settings');document.getElementById('baseUrl').value=s.public_base_url||'';document.getElementById('apiKey').placeholder=s.gemini_key_configured?'API key already saved (enter a new key to replace)':'Paste Gemini API key'}catch(e){showMessage(e.message)}}
async function saveSettings(){try{let p={public_base_url:document.getElementById('baseUrl').value};let k=document.getElementById('apiKey').value;if(k)p.gemini_api_key=k;await api('/api/settings',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify(p)});document.getElementById('apiKey').value='';showMessage('Settings saved.')}catch(e){showMessage(e.message)}}
async function loadModels(){try{let k=document.getElementById('apiKey').value;let d=await api('/api/gemini/models',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({api_key:k})});document.getElementById('models').textContent='Available models: '+d.models.map(x=>x.name).join(', ')}catch(e){document.getElementById('models').textContent=e.message}}initSettings();'''
    return render_page("Settings", "settings", content, script)


@app.get("/api-guide", response_class=HTMLResponse)
def api_guide_page():
    content = '''<h1>AI API Guide</h1><p class="muted">AI Agent বা অন্য software-কে product data খোঁজার জন্য এই endpoints ব্যবহার করতে দাও।</p>
<section class="card"><h2>1. Full JSON data</h2><pre>GET /api/public/{client_id}</pre><p class="small">সম্পূর্ণ product list ফেরত দেয়। বড় তালিকায় সবসময় এটি ব্যবহার করার প্রয়োজন নেই।</p></section>
<section class="card"><h2>2. Universal Search & Filter</h2><pre>GET /api/public/{client_id}/search?q=note
GET /api/public/{client_id}/search?q=tank
GET /api/public/{client_id}/search?brand=XIAOMI&min_price=10000&max_price=30000
GET /api/public/{client_id}/search?q=note&sort=price_asc&page=1&page_size=20</pre>
<p>Supported parameters:</p><ul><li><b>q</b>: brand, model, variant, notes-এর মধ্যে exact substring search (case-insensitive)।</li><li><b>brand</b>: নির্দিষ্ট brand filter।</li><li><b>min_price / max_price</b>: price range।</li><li><b>sort</b>: price_asc, price_desc, name_asc, name_desc।</li><li><b>page / page_size</b>: pagination।</li></ul><p class="small">কোনো fuzzy বা কাছাকাছি মডেল নিজে থেকে যোগ করা হয় না। Search শব্দের সঙ্গে সরাসরি মেলা রেকর্ডই ফেরত আসে।</p></section>
<section class="card"><h2>3. Available brands</h2><pre>GET /api/public/{client_id}/brands</pre></section>
<section class="card"><h2>4. AI Tool Function Schema</h2><pre>GET /api/public/{client_id}/tool-schema</pre><p class="small">এটি AI tool/function integration-এর জন্য নাম, description ও JSON parameter schema দেয়। AI platform-এ tool register করার সময় schema ব্যবহার করতে হবে।</p></section>
<section class="card"><h2>5. Other links</h2><pre>GET /data/{client_id}                 # Public HTML page
GET /download/{client_id}.xlsx        # Excel download
GET /api/clients/{client_id}/changes  # Recent processing history</pre></section>'''
    return render_page("API Guide", "api", content)


@app.get("/health")
def health():
    return {"status": "ok", "app": APP_TITLE, "time": now_iso()}


@app.get("/api/settings")
def get_settings():
    key = setting_get("gemini_api_key")
    return {
        "gemini_key_configured": bool(key),
        "public_base_url": setting_get("public_base_url"),
        "poll_minutes": POLL_MINUTES,
    }


@app.post("/api/settings")
def save_settings(payload: dict):
    if "gemini_api_key" in payload:
        setting_set("gemini_api_key", str(payload["gemini_api_key"]).strip())
    if "public_base_url" in payload:
        setting_set("public_base_url", str(payload["public_base_url"]).strip().rstrip("/"))
    return {"ok": True, "message": "Settings saved"}


@app.post("/api/gemini/models")
def gemini_models(payload: dict):
    key = str(payload.get("api_key") or setting_get("gemini_api_key")).strip()
    if not key:
        raise HTTPException(400, "Gemini API Key দাও।")
    try:
        return {"models": discover_gemini_models(key)}
    except Exception as exc:
        raise HTTPException(400, str(exc))


@app.get("/api/clients")
def list_clients(request: Request):
    with db() as con:
        rows = con.execute("SELECT * FROM clients ORDER BY created_at DESC").fetchall()
    base = public_base_url(request)
    result = []
    for row in rows:
        c = dict(row)
        products = json.loads(c.get("products_json") or "[]")
        result.append({
            "id": c["id"], "name": c["name"], "source_url": c["source_url"],
            "source_type": c["source_type"], "total_products": len(products),
            "status": c["status"], "last_checked": c["last_checked"],
            "last_updated": c["last_updated"], "last_error": c["last_error"],
            "public_url": f"{base}/data/{c['id']}" if base else f"/data/{c['id']}",
            "api_url": f"{base}/api/public/{c['id']}" if base else f"/api/public/{c['id']}",
        })
    return result


@app.post("/api/clients")
def create_client(payload: dict):
    name = str(payload.get("name") or "").strip()
    source_url = str(payload.get("source_url") or "").strip()
    if not name:
        raise HTTPException(400, "Client name প্রয়োজন।")
    if source_url and "docs.google.com/document" not in source_url:
        raise HTTPException(400, "এখানে Google Docs document URL দিতে হবে।")
    client_id = make_client_id(name)
    source_type = "google_docs" if source_url else "text"
    with db() as con:
        con.execute(
            "INSERT INTO clients(id,name,source_url,source_type,created_at) VALUES(?,?,?,?,?)",
            (client_id, name, source_url, source_type, now_iso()),
        )
        con.commit()
    log_event(client_id, "Created", "Client profile created")
    return {"ok": True, "id": client_id, "name": name}


@app.post("/api/clients/{client_id}/source")
def update_source_url(client_id: str, payload: dict):
    get_client(client_id)
    source_url = str(payload.get("source_url") or "").strip()
    if source_url and "docs.google.com/document" not in source_url:
        raise HTTPException(400, "এখানে Google Docs document URL দিতে হবে।")
    update_client(client_id, source_url=source_url, source_type="google_docs" if source_url else "text", source_text="")
    log_event(client_id, "Source URL updated", source_url or "Source URL cleared")
    return {"ok": True, "message": "Source URL saved. Process Now চাপলে ডেটা আপডেট হবে।"}


@app.post("/api/clients/{client_id}/upload")
async def upload_source(client_id: str, file: UploadFile = File(...)):
    get_client(client_id)
    raw = await file.read()
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "ফাইল সর্বোচ্চ 15 MB হতে পারবে।")
    try:
        text = read_uploaded_file(file.filename or "upload.txt", raw)
    except Exception as exc:
        raise HTTPException(400, str(exc))
    if not text.strip():
        raise HTTPException(400, "ফাইলে কোনো টেক্সট পাওয়া যায়নি।")
    update_client(client_id, source_type="upload", source_text=text, source_url="")
    log_event(client_id, "Source uploaded", file.filename or "file")
    return {"ok": True, "characters": len(text), "message": "Source uploaded. এখন Process চাপো।"}


@app.post("/api/clients/{client_id}/process")
def process_now(client_id: str):
    return process_client(client_id, force=True)


@app.post("/api/clients/{client_id}/check")
def check_now(client_id: str):
    return process_client(client_id, force=False)


@app.get("/api/clients/{client_id}")
def client_detail(client_id: str):
    c = get_client(client_id)
    products = json.loads(c.get("products_json") or "[]")
    with db() as con:
        history = con.execute(
            "SELECT event,details,created_at FROM history WHERE client_id=? ORDER BY id DESC LIMIT 30",
            (client_id,),
        ).fetchall()
    return {
        "id": c["id"], "name": c["name"], "source_url": c["source_url"],
        "source_type": c["source_type"], "products": products,
        "status": c["status"], "last_checked": c["last_checked"],
        "last_updated": c["last_updated"], "last_error": c["last_error"],
        "history": [dict(x) for x in history],
    }


@app.delete("/api/clients/{client_id}")
def delete_client(client_id: str):
    get_client(client_id)
    with db() as con:
        con.execute("DELETE FROM clients WHERE id=?", (client_id,))
        con.execute("DELETE FROM history WHERE client_id=?", (client_id,))
        con.commit()
    return {"ok": True}


@app.get("/api/public/{client_id}")
def public_json(client_id: str):
    c = get_client(client_id)
    products = json.loads(c.get("products_json") or "[]")
    return JSONResponse({
        "client_id": c["id"],
        "client_name": c["name"],
        "last_updated": c["last_updated"],
        "total_products": len(products),
        "products": products,
    })


def product_matches_query(product, query):
    q = str(query or "").strip().casefold()
    if not q:
        return True
    searchable = " ".join(str(product.get(k) or "") for k in ("brand", "model", "variant", "notes"))
    return q in searchable.casefold()


@app.get("/api/public/{client_id}/brands")
def public_brands(client_id: str):
    c = get_client(client_id)
    products = json.loads(c.get("products_json") or "[]")
    brands = sorted({str(p.get("brand") or "").strip() for p in products if str(p.get("brand") or "").strip()}, key=str.casefold)
    return {"client_id": client_id, "total_brands": len(brands), "brands": brands}


@app.get("/api/public/{client_id}/search")
def public_search(
    client_id: str,
    q: str = "",
    brand: str = "",
    min_price: float | None = None,
    max_price: float | None = None,
    sort: str = "price_asc",
    page: int = 1,
    page_size: int = 50,
):
    c = get_client(client_id)
    products = json.loads(c.get("products_json") or "[]")
    if min_price is not None and min_price < 0:
        raise HTTPException(400, "min_price ঋণাত্মক হতে পারে না।")
    if max_price is not None and max_price < 0:
        raise HTTPException(400, "max_price ঋণাত্মক হতে পারে না।")
    if min_price is not None and max_price is not None and min_price > max_price:
        raise HTTPException(400, "min_price, max_price-এর চেয়ে বড় হতে পারবে না।")
    page = max(1, page)
    page_size = min(200, max(1, page_size))
    filtered = []
    for product in products:
        if not product_matches_query(product, q):
            continue
        pbrand = str(product.get("brand") or "")
        if brand and pbrand.casefold() != brand.strip().casefold():
            continue
        price = product.get("price")
        try:
            numeric_price = float(price) if price not in (None, "") else None
        except (TypeError, ValueError):
            numeric_price = None
        if min_price is not None and (numeric_price is None or numeric_price < min_price):
            continue
        if max_price is not None and (numeric_price is None or numeric_price > max_price):
            continue
        filtered.append(product)

    sort = sort.lower().strip()
    if sort not in {"price_asc", "price_desc", "name_asc", "name_desc"}:
        raise HTTPException(400, "sort অবশ্যই price_asc, price_desc, name_asc অথবা name_desc হবে।")
    def price_key(p):
        try:
            return (p.get("price") in (None, ""), float(p.get("price")) if p.get("price") not in (None, "") else float("inf"))
        except (TypeError, ValueError):
            return (True, float("inf"))
    if sort == "price_asc":
        filtered.sort(key=price_key)
    elif sort == "price_desc":
        filtered.sort(key=lambda p: (price_key(p)[0], -price_key(p)[1] if not price_key(p)[0] else 0))
    elif sort == "name_asc":
        filtered.sort(key=lambda p: (str(p.get("brand") or "").casefold(), str(p.get("model") or "").casefold()))
    else:
        filtered.sort(key=lambda p: (str(p.get("brand") or "").casefold(), str(p.get("model") or "").casefold()), reverse=True)
    total = len(filtered)
    start = (page - 1) * page_size
    page_items = filtered[start:start + page_size]
    return {
        "client_id": c["id"], "client_name": c["name"], "query": q,
        "filters": {"brand": brand or None, "min_price": min_price, "max_price": max_price, "sort": sort},
        "total_matches": total, "page": page, "page_size": page_size,
        "total_pages": (total + page_size - 1) // page_size,
        "has_more": start + page_size < total,
        "products": page_items,
    }


@app.get("/api/public/{client_id}/tool-schema")
def public_tool_schema(client_id: str):
    get_client(client_id)
    return {
        "type": "function", "name": "search_client_products",
        "description": "Search only products that exist in this client's price list. Searches brand, model, variant and notes using direct case-insensitive substring matching. Does not invent or return approximate products.",
        "parameters": {
            "type": "object", "properties": {
                "q": {"type": "string", "description": "Any brand, model, series, variant or notes keyword. Example: Note or Tank."},
                "brand": {"type": "string", "description": "Optional exact brand filter."},
                "min_price": {"type": "number"}, "max_price": {"type": "number"},
                "sort": {"type": "string", "enum": ["price_asc", "price_desc", "name_asc", "name_desc"]},
                "page": {"type": "integer", "minimum": 1},
                "page_size": {"type": "integer", "minimum": 1, "maximum": 200}
            }, "required": []
        },
        "endpoint": f"/api/public/{client_id}/search", "method": "GET"
    }


@app.get("/data/{client_id}", response_class=HTMLResponse)
def public_view(client_id: str):
    c = get_client(client_id)
    products = json.loads(c.get("products_json") or "[]")
    safe_name = (c["name"].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))

    # Sort priced products from lowest to highest; products without a valid price go last.
    def price_sort_key(product):
        price = product.get("price")
        try:
            return (price is None or price == "", float(price) if price not in (None, "") else float("inf"))
        except (TypeError, ValueError):
            return (True, float("inf"))

    sorted_products = sorted(products, key=price_sort_key)
    brands = sorted(
        {str(p.get("brand") or "").strip() for p in products if str(p.get("brand") or "").strip()},
        key=str.casefold,
    )
    priced_values = []
    for product in products:
        try:
            if product.get("price") not in (None, ""):
                priced_values.append(float(product["price"]))
        except (TypeError, ValueError):
            pass
    min_price = min(priced_values) if priced_values else None
    max_price = max(priced_values) if priced_values else None

    rows = []
    for product in sorted_products:
        brand = str(product.get("brand") or "")
        model = str(product.get("model") or "")
        variant = str(product.get("variant") or "")
        currency = str(product.get("currency") or "")
        notes = str(product.get("notes") or "")
        price = product.get("price")
        try:
            numeric_price = float(price) if price not in (None, "") else ""
        except (TypeError, ValueError):
            numeric_price = ""
        values = [brand, model, variant, price, currency, notes]
        cells = "".join(
            "<td>" + str(value if value is not None else "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;") + "</td>"
            for value in values
        )
        brand_attr = brand.replace("&", "&amp;").replace('"', "&quot;").replace("<", "&lt;").replace(">", "&gt;")
        search_text = " ".join(str(value if value is not None else "") for value in values).replace("&", "&amp;").replace('"', "&quot;").replace("<", "&lt;").replace(">", "&gt;")
        rows.append(
            f'<tr data-brand="{brand_attr}" data-price="{numeric_price}" data-search="{search_text.lower()}">{cells}</tr>'
        )
    table = "".join(rows) or '<tr><td colspan="6">No processed products yet.</td></tr>'
    brand_options = ''.join(
        '<option value="' + b.replace("&", "&amp;").replace('"', "&quot;").replace("<", "&lt;").replace(">", "&gt;") + '">' +
        b.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;") + '</option>'
        for b in brands
    )
    min_label = f'{min_price:,.0f}' if min_price is not None else 'N/A'
    max_label = f'{max_price:,.0f}' if max_price is not None else 'N/A'
    html = f'''<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
    <title>{safe_name} - Price List</title><style>
    *{{box-sizing:border-box}}body{{font-family:Arial,sans-serif;background:#f4f6fa;margin:0;color:#172033}}
    header{{background:#13213b;color:white;padding:24px}}main{{max-width:1200px;margin:24px auto;padding:0 16px}}
    .meta{{color:#667085;font-size:13px;margin:12px 0}}.wrap{{overflow:auto;background:white;border-radius:14px;box-shadow:0 4px 22px #17203312}}
    table{{border-collapse:collapse;width:100%;min-width:700px}}th,td{{padding:12px 14px;border-bottom:1px solid #edf0f4;text-align:left}}
    th{{background:#eaf0f8;position:sticky;top:0;white-space:nowrap}}tr:hover td{{background:#f8fbff}}
    a{{color:#1d63c4;text-decoration:none}}.top{{display:flex;gap:16px;flex-wrap:wrap;justify-content:space-between;align-items:center}}
    .filters{{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px;background:white;padding:16px;border-radius:14px;box-shadow:0 4px 22px #1720330c}}
    .field{{display:flex;flex-direction:column;gap:6px}}label{{font-size:12px;font-weight:700;color:#475467}}
    input,select{{width:100%;padding:11px 12px;border:1px solid #d5dbe5;border-radius:8px;background:white;font-size:14px;color:#172033}}
    button{{padding:11px 15px;border:0;border-radius:8px;background:#173b70;color:white;font-weight:700;cursor:pointer;align-self:end}}
    .count{{font-size:13px;color:#475467;margin:12px 0}}@media(max-width:600px){{header{{padding:18px}}main{{margin:16px auto}}}}
    </style></head><body><header><div class="top"><div><h1>{safe_name}</h1><div>Public Product Price List</div></div>
    <div><a style="color:white" href="/api/public/{client_id}">JSON API</a> · <a style="color:white" href="/download/{client_id}.xlsx">Download Excel</a></div></div></header>
    <main>
      <section class="filters" aria-label="Product filters">
        <div class="field"><label for="q">SEARCH PRODUCTS</label><input id="q" placeholder="Search brand, model, variant..." oninput="filterRows()"></div>
        <div class="field"><label for="brandFilter">BRAND</label><select id="brandFilter" onchange="filterRows()"><option value="">All brands</option>{brand_options}</select></div>
        <div class="field"><label for="minPrice">MINIMUM PRICE (BDT)</label><input id="minPrice" type="number" min="0" step="any" placeholder="৳{min_label}" oninput="filterRows()"></div>
        <div class="field"><label for="maxPrice">MAXIMUM PRICE (BDT)</label><input id="maxPrice" type="number" min="0" step="any" placeholder="৳{max_label}" oninput="filterRows()"></div>
        <button type="button" onclick="resetFilters()">Reset filters</button>
      </section>
      <div class="meta">Total products: {len(products)} · Price range: ৳{min_label} – ৳{max_label} · Default order: lowest to highest · Last updated: {c.get('last_updated') or 'Not processed yet'}</div>
      <div class="count" id="resultCount">Showing {len(products)} products</div>
      <div class="wrap"><table id="tbl"><thead><tr><th>Brand</th><th>Model</th><th>Variant</th><th>Price</th><th>Currency</th><th>Notes</th></tr></thead><tbody>{table}</tbody></table></div>
    </main><script>
    function filterRows(){{
      const q=document.getElementById('q').value.trim().toLowerCase();
      const brand=document.getElementById('brandFilter').value;
      const minRaw=document.getElementById('minPrice').value;
      const maxRaw=document.getElementById('maxPrice').value;
      const min=minRaw===''?null:Number(minRaw), max=maxRaw===''?null:Number(maxRaw);
      let shown=0;
      document.querySelectorAll('#tbl tbody tr').forEach(row=>{{
        const rowBrand=row.dataset.brand||'', rawPrice=row.dataset.price;
        const price=rawPrice===''?null:Number(rawPrice);
        const matchesSearch=(row.dataset.search||'').includes(q);
        const matchesBrand=!brand||rowBrand.toLowerCase()===brand.toLowerCase();
        const matchesMin=min===null||(price!==null&&price>=min);
        const matchesMax=max===null||(price!==null&&price<=max);
        const visible=matchesSearch&&matchesBrand&&matchesMin&&matchesMax;
        row.style.display=visible?'':'none';if(visible)shown++;
      }});
      document.getElementById('resultCount').textContent='Showing '+shown+' of {len(products)} products';
    }}
    function resetFilters(){{document.getElementById('q').value='';document.getElementById('brandFilter').value='';document.getElementById('minPrice').value='';document.getElementById('maxPrice').value='';filterRows();}}
    </script></body></html>'''
    return HTMLResponse(html)


@app.get("/download/{client_id}.xlsx")
def download_excel(client_id: str):
    c = get_client(client_id)
    products = json.loads(c.get("products_json") or "[]")
    stream = build_excel(products, c["name"])
    filename = re.sub(r"[^A-Za-z0-9_-]+", "_", c["name"])[:50] or "price-list"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}.xlsx"'},
    )


@app.get("/api/clients/{client_id}/changes")
def client_changes(client_id: str):
    get_client(client_id)
    with db() as con:
        rows = con.execute(
            "SELECT event,details,created_at FROM history WHERE client_id=? ORDER BY id DESC LIMIT 20",
            (client_id,),
        ).fetchall()
    return {"client_id": client_id, "events": [dict(r) for r in rows]}


